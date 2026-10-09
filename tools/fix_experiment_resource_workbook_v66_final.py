#!/usr/bin/env python3
"""Use fixed preregistered z-quantiles to avoid locale-specific Excel UDF names."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "实验资源与样本量预算v6.6.xlsx"
Z_ALPHA = 1.95996398454005
Z_POWER = 1.2815515655446


def main() -> None:
    wb = load_workbook(OUT)
    ws = wb.worksheets[3]
    for row in range(4, 16):
        ws.cell(row, 6, f"=ROUNDUP((({Z_ALPHA}+{Z_POWER})^2*B{row})/(C{row}^2),0)")
    for row in range(16, 20):
        ws.cell(row, 6, f"=ROUNDUP((({Z_ALPHA}+{Z_POWER})/B{row})^2,0)")
    ws["G3"] = "解释（z固定为单侧alpha=.025、power=.90的标准正态分位数）"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

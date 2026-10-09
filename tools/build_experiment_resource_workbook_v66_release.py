#!/usr/bin/env python3
"""Release builder for the v6.6 resource workbook."""

from __future__ import annotations

import importlib.util
from itertools import product
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "tools" / "build_experiment_resource_workbook_v66.py"
OUT = ROOT / "实验资源与样本量预算v6.6.xlsx"


def load_base():
    spec = importlib.util.spec_from_file_location("v66_base", BASE)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load base workbook builder")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def install_power_fix(module) -> None:
    def build_power(wb):
        ws = wb.create_sheet("功效敏感性")
        module.title(ws, "确认性样本量敏感性（仅用于资源上限，不是正式功效冻结）", 7)
        module.header(ws, 3, ["类型", "配对不一致率 q / 标准化效应 dz", "规划真值越过 margin 的间隔 g", "单侧 alpha", "目标 power", "近似 n", "解释"])
        row = 4
        for q, gap in product([0.10, 0.20, 0.30], [0.02, 0.03, 0.04, 0.05]):
            values = ["paired_binary", q, gap, 0.025, 0.90, None,
                      "McNemar/配对风险差的保守正态近似；g 是规划真值超过 SESOI/no-harm margin 的额外距离，不是 SESOI 本身。"]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            ws.cell(row, 6, f"=ROUNDUP(((NORM.S.INV(1-D{row})+NORM.S.INV(E{row}))^2*B{row})/(C{row}^2),0)")
            for col in [2, 3, 4, 5]:
                module.input_cell(ws.cell(row, col), "Planning sensitivity grid; replace with held-out pilot kernel and simulation-based power.")
                ws.cell(row, col).number_format = "0.0%"
            module.formula_cell(ws.cell(row, 6))
            row += 1
        for dz in [0.10, 0.15, 0.20, 0.25]:
            values = ["paired_delay", dz, "-", 0.025, 0.90, None,
                      "标准化配对差近似；正式分析必须模拟截尾、miss mass、分层权重、variant/batch 相关和联合 IUT。"]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            ws.cell(row, 6, f"=ROUNDUP(((NORM.S.INV(1-D{row})+NORM.S.INV(E{row}))/B{row})^2,0)")
            for col in [2, 4, 5]:
                module.input_cell(ws.cell(row, col), "Planning sensitivity grid only.")
            ws.cell(row, 2).number_format = "0.00"
            ws.cell(row, 4).number_format = ws.cell(row, 5).number_format = "0.0%"
            module.formula_cell(ws.cell(row, 6))
            row += 1
        module.style_table(ws, 4, row - 1, 1, 7)
        chart = LineChart()
        chart.title = "配对二元终点：n 对 margin 越界间隔的敏感性"
        chart.y_axis.title = "近似 n"
        chart.x_axis.title = "敏感性网格行（按 q 分组）"
        chart.add_data(Reference(ws, min_col=6, min_row=3, max_row=15), titles_from_data=True)
        chart.height = 7
        chart.width = 14
        ws.add_chart(chart, "I4")
        module.autosize(ws, {1: 18, 2: 23, 3: 25, 4: 14, 5: 14, 6: 14, 7: 72})

    module.build_power = build_power


def apply_capacity_envelope() -> None:
    wb = load_workbook(OUT)
    readme = wb["说明"]
    readme["B6"] = "32 逻辑线程、128 GB RAM、2 TB 可用 NVMe、1×NVIDIA T4"
    readme["C6"] = "约 169 人日"
    readme["D6"] = "约 120 T4 GPU-h；约 1.2k CPU-h 中央估计。"
    readme["B7"] = "功效初始工作域约2,500–10,000个可分析bundle；资源按14,400预留"
    readme["C7"] = "推荐容量点 n_H/n_G/n_B=9,600/2,400/2,400"
    readme["D7"] = "正式 n 只允许在独立 pilot 与 MC confirm 后冻结；n_max_generate 建议24,000。"
    inputs = wb["假设输入"]
    replacements = {
        "hazard_confirm_bundles": (4800, 9600, 14400),
        "benign_confirm_bundles": (1200, 2400, 3600),
        "budget_confirm_bundles": (1200, 2400, 3600),
        "hazard_methods": (4, 5, 6),
        "benign_methods": (2, 3, 4),
        "budget_methods": (2, 5, 6),
        "variants": (1, 3, 5),
        "algorithm_seeds": (1, 2, 3),
        "pretest_method_runs": (12000, 144000, 864000),
        "external_method_runs": (3600, 90000, 720000),
        "power_sim_cpu_hours": (50, 250, 1000),
        "retained_mb_per_method_run": (0.1, 0.3, 0.5),
        "gpu_hours": (0, 120, 300),
    }
    for row in range(4, inputs.max_row + 1):
        name = inputs.cell(row, 1).value
        if name in replacements:
            for col, value in zip(range(2, 5), replacements[name]):
                inputs.cell(row, col, value)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)


def main() -> None:
    module = load_base()
    install_power_fix(module)
    module.main()
    apply_capacity_envelope()
    print(OUT)


if __name__ == "__main__":
    main()

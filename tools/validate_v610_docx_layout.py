from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "实验设计与执行计划v6.10-100GB分阶段执行版.docx"
OUT = ROOT / "v6.10-DOCX布局与宽度核验.json"


def main() -> None:
    doc = Document(DOCX)
    failures: list[str] = []
    section_checks: list[dict[str, object]] = []
    minimum_content_width = None
    for index, section in enumerate(doc.sections, start=1):
        content_width = section.page_width - section.left_margin - section.right_margin
        minimum_content_width = content_width if minimum_content_width is None else min(minimum_content_width, content_width)
        section_checks.append(
            {
                "section": index,
                "page_width_twips": section.page_width.twips,
                "page_height_twips": section.page_height.twips,
                "left_margin_twips": section.left_margin.twips,
                "right_margin_twips": section.right_margin.twips,
                "content_width_twips": content_width.twips,
            }
        )

    table_checks: list[dict[str, object]] = []
    assert minimum_content_width is not None
    for index, table in enumerate(doc.tables, start=1):
        grid = table._tbl.tblGrid
        widths: list[int] = []
        if grid is not None:
            for column in grid.gridCol_lst:
                value = column.get(qn("w:w"))
                if value:
                    widths.append(int(value))
        grid_total = sum(widths)
        overflow = max(0, grid_total - minimum_content_width.twips)
        if not table.rows or not table.columns:
            failures.append(f"table {index} has no rows or columns")
        if grid_total and overflow > 20:
            failures.append(f"table {index} grid exceeds content width by {overflow} twips")
        table_checks.append(
            {
                "table": index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "grid_widths_twips": widths,
                "grid_total_twips": grid_total,
                "content_width_twips": minimum_content_width.twips,
                "overflow_twips": overflow,
            }
        )

    with zipfile.ZipFile(DOCX) as archive:
        zip_failure = archive.testzip()
        xml = archive.read("word/document.xml").decode("utf-8")
    if zip_failure is not None:
        failures.append(f"ZIP integrity failure: {zip_failure}")
    if "100,000,000,000" not in "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells):
        failures.append("100GB contract token missing")
    if "<w:ins" in xml or "<w:del" in xml or "<w:commentRangeStart" in xml:
        failures.append("tracked changes or comments remain")

    report = {
        "status": "STRUCTURAL_LAYOUT_PASS_RENDER_DEFERRED" if not failures else "FAIL",
        "render_status": "DEFERRED_ACTIVE_WORD_INSTANCE",
        "reason": "An active Word process prevented an isolated PDF export; no user process was terminated.",
        "sections": section_checks,
        "tables": table_checks,
        "paragraphs": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "all_table_grids_within_content_width": not any(check["overflow_twips"] > 20 for check in table_checks),
        "zip_test": zip_failure,
        "failures": failures,
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(OUT), "status": report["status"], "tables": len(table_checks), "failures": failures}, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

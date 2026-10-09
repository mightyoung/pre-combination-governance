#!/usr/bin/env python3
"""Build the v6.6 Word reading copy when docx-js is unavailable."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "实验启动计划v6.6-双机执行与资源预算.md"
OUT = ROOT / "实验启动计划v6.6-双机执行与资源预算.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    paragraph.add_run(" 页")


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(127, 96, 0)
        else:
            run = paragraph.add_run(token)
            run.font.color.rgb = RGBColor(5, 99, 193)
            run.font.underline = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and re.match(r"^\s*\|.*\|\s*$", lines[i]):
        cells = [x.strip() for x in lines[i].strip()[1:-1].split("|")]
        if not all(re.match(r"^:?-{3,}:?$", x) for x in cells):
            rows.append(cells)
        i += 1
    return rows, i


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, color in [
        ("Title", 20, RGBColor(23, 54, 93)),
        ("Heading 1", 15, RGBColor(23, 54, 93)),
        ("Heading 2", 13, RGBColor(47, 85, 151)),
        ("Heading 3", 11.5, RGBColor(68, 84, 106)),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 3"].paragraph_format.space_before = Pt(6)

    if "Code Block" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor(127, 96, 0)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for r, values in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, values[c] if c < len(values) else "")
            for run in p.runs:
                run.font.name = "Microsoft YaHei"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                run.font.size = Pt(8.5)
                if r == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r == 0:
                set_cell_shading(cell, "4472C4")
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    header = section.header.paragraphs[0]
    header.text = "实验启动计划 v6.6 | PLANNING-ONLY"
    header.style = doc.styles["Normal"]
    header.runs[0].font.color.rgb = RGBColor(91, 101, 115)
    header.runs[0].font.size = Pt(9)
    add_page_number(section.footer.paragraphs[0])

    lines = SRC.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    i = 0
    title_seen = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*:?-{3,}", lines[i + 1]):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Title")
                add_inline(p, text)
                subtitle = doc.add_paragraph()
                subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = subtitle.add_run("基于 v20 / v6.5 | 2026-07-14")
                run.italic = True
                run.font.color.rgb = RGBColor(91, 101, 115)
                toc_heading = doc.add_paragraph("目录", style="Heading 1")
                toc_heading.paragraph_format.page_break_before = False
                add_toc(doc.add_paragraph())
                title_seen = True
            else:
                p = doc.add_paragraph(style=f"Heading {level - 1}")
                add_inline(p, text)
                if level == 2 and title_seen and not text.startswith("Material Passport"):
                    p.paragraph_format.page_break_before = True
            i += 1
            continue
        bullet = re.match(r"^\s*-\s+(.*)$", raw)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", raw)
        if bullet or numbered:
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(p, (bullet or numbered).group(1))
            i += 1
            continue
        if (line.startswith("`") and line.endswith("`。")) or line.startswith("n ≈"):
            p = doc.add_paragraph(style="Code Block")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line.replace("`", ""))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, line)
        i += 1

    doc.core_properties.title = "实验启动计划 v6.6：双机执行、功效路线与资源预算"
    doc.core_properties.subject = "v20/v6.5 experiment execution and resource plan"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

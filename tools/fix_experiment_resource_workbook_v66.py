#!/usr/bin/env python3
"""Normalize formulas for broad Excel compatibility in the v6.6 workbook."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "实验资源与样本量预算v6.6.xlsx"


def main() -> None:
    wb = load_workbook(OUT)
    ws = wb.worksheets[3]
    for row in range(4, 20):
        cell = ws.cell(row, 6)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.value = cell.value.replace("NORM.S.INV", "NORMSINV")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

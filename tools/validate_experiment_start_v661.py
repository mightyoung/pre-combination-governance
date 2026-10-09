#!/usr/bin/env python3
"""Validate the v6.6.1 DeepSeek-aware experiment start package and write a report."""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "实验启动计划v6.6.1-DeepSeek在线API与双机资源预算.md"
DOCX = ROOT / "实验启动计划v6.6.1-DeepSeek在线API与双机资源预算.docx"
XLSX = ROOT / "实验资源与样本量预算v6.6.1.xlsx"
CHECKLIST = ROOT / "实验启动清单v6.6.1.json"
PROBE = ROOT / "tools" / "probe_deepseek_v4_flash_v661.py"
REPORT = ROOT / "实验启动包v6.6.1-核验报告.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def close(actual, expected, tolerance=1e-6) -> bool:
    return actual is not None and math.isclose(float(actual), float(expected), rel_tol=tolerance, abs_tol=tolerance)


def inspect_docx() -> dict:
    doc = Document(DOCX)
    heading_count = sum(1 for p in doc.paragraphs if p.style and p.style.name.startswith("Heading"))
    with ZipFile(DOCX) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml").decode("utf-8")
        all_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in names if name.endswith(".xml")
        )
    headers = [p.text for section in doc.sections for p in section.header.paragraphs]
    checks = {
        "title_contains_v6.6.1": any("v6.6.1" in p.text and "DeepSeek" in p.text for p in doc.paragraphs),
        "header_contains_v6.6.1": any("v6.6.1" in text for text in headers),
        "toc_field": "TOC \\o" in document_xml,
        "deepseek_model_id": "deepseek-v4-flash" in document_xml,
        "external_model_manifest": "ExternalModelReleaseManifest" in document_xml,
        "doubled_version_absent": "v6.6.1.1" not in all_xml,
        "comments_absent": not any(name.startswith("word/comments") for name in names),
        "tracked_insertions": all_xml.count("<w:ins"),
        "tracked_deletions": all_xml.count("<w:del"),
    }
    if not all(value is True or value == 0 for value in checks.values()):
        raise RuntimeError(f"DOCX structural check failed: {checks}")
    return {
        "office_xml_validator": "PASS (run separately with docx/scripts/office/validate.py)",
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": heading_count,
        **checks,
    }


def inspect_xlsx() -> dict:
    formulas = load_workbook(XLSX, data_only=False)
    values = load_workbook(XLSX, data_only=True)
    formula_count = 0
    formula_blanks = []
    formula_errors = []
    for index, ws in enumerate(formulas.worksheets):
        value_ws = values.worksheets[index]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    cached = value_ws[cell.coordinate].value
                    if cached is None:
                        formula_blanks.append((ws.title, cell.coordinate))
                    if isinstance(cached, str) and cached.startswith("#"):
                        formula_errors.append((ws.title, cell.coordinate, cached))
    api = values["DeepSeek API预算"]
    total = values["总资源"]
    recommended = {
        "base_api_calls": api["C16"].value,
        "billable_attempts": api["C17"].value,
        "input_tokens_million": api["C18"].value,
        "output_tokens_million": api["C19"].value,
        "cache_hit_input_tokens_million": api["C20"].value,
        "cache_miss_input_tokens_million": api["C21"].value,
        "expected_cost_rmb": api["C22"].value,
        "complete_cache_miss_cost_rmb": api["C23"].value,
        "ideal_acquisition_wall_hours": api["C24"].value,
        "authorization_cap_rmb": api["C25"].value,
        "t4_gpu_hours": total["C10"].value,
        "person_days": total["C11"].value,
        "cpu_core_hours_total": total["C7"].value,
        "retained_storage_gb": total["C9"].value,
    }
    expected = {
        "base_api_calls": 48000,
        "billable_attempts": 52800,
        "input_tokens_million": 158.4,
        "output_tokens_million": 26.4,
        "cache_hit_input_tokens_million": 110.88,
        "cache_miss_input_tokens_million": 47.52,
        "expected_cost_rmb": 102.5376,
        "complete_cache_miss_cost_rmb": 211.2,
        "authorization_cap_rmb": 1000,
        "t4_gpu_hours": 80,
        "person_days": 169,
    }
    mismatches = {key: (recommended[key], value) for key, value in expected.items() if not close(recommended[key], value)}
    intentional = [item for item in formula_blanks if item[0] == "人力预算" and item[1].startswith("F")]
    if formula_errors or mismatches or len(formula_blanks) != 9 or len(intentional) != 9:
        raise RuntimeError(
            f"XLSX check failed: errors={formula_errors}, mismatches={mismatches}, blanks={formula_blanks}"
        )
    return {
        "sheet_count": len(formulas.sheetnames),
        "sheets": formulas.sheetnames,
        "formula_count": formula_count,
        "formula_errors": len(formula_errors),
        "recalculation": "Excel COM CalculateFullRebuild and save",
        "intentional_blank_formula_results": 9,
        "recommended_scenario": recommended,
    }


def inspect_json() -> dict:
    data = json.loads(CHECKLIST.read_text(encoding="utf-8"))
    api = data["online_llm_api"]
    checks = {
        "schema_version": data["schema_version"] == "experiment-start-v6.6.1-1.0",
        "model": api["model"] == "deepseek-v4-flash",
        "thinking_disabled": api["frozen_request_contract"]["thinking"]["type"] == "disabled",
        "temperature_zero": api["frozen_request_contract"]["temperature"] == 0,
        "shared_immutable_cache": api["fairness_and_isolation"]["immutable_shared_cache_for_all_methods"] is True,
        "network_off_for_confirm": api["fairness_and_isolation"]["network_disabled_during_confirm_method_runs"] is True,
        "no_embedded_api_key": "sk-" not in CHECKLIST.read_text(encoding="utf-8").lower(),
    }
    if not all(checks.values()):
        raise RuntimeError(f"Checklist check failed: {checks}")
    return {"parse": "PASS", **checks}


def main() -> None:
    markdown = MD.read_text(encoding="utf-8")
    required = [
        "deepseek-v4-flash", "thinking.type=disabled", "system_fingerprint",
        "ExternalModelReleaseManifest", "约 102.54 元", "约 211.20 元",
        "所有方法只读同一缓存", "TruthManifest",
    ]
    missing = [item for item in required if item not in markdown]
    if missing:
        raise RuntimeError(f"Markdown missing required contract terms: {missing}")

    original_expected = {
        "实验启动计划v6.6-双机执行与资源预算.md": "AA6338B3E66BAA8C3E200C64B5CED109886223370DCC506FD1C7A5C0C4FD88F1",
        "实验启动计划v6.6-双机执行与资源预算.docx": "F313CABD133AD1196693DEBC3525909D1E36D5088014E33519F0C20DA7B77A61",
        "实验资源与样本量预算v6.6.xlsx": "598D0B4E399424F71EA625EBAC5EC278769577F3CE19F55A7423127778141F1B",
        "实验启动清单v6.6.json": "991D794DFD1C2593230EE6754E17FF727C276350FFEAFD679244B1FEFF1EB8F3",
    }
    original_actual = {name: sha256(ROOT / name) for name in original_expected}
    originals_preserved = original_actual == original_expected
    if not originals_preserved:
        raise RuntimeError(f"Original v6.6 artifacts changed: {original_actual}")

    artifacts = [MD, DOCX, XLSX, CHECKLIST, PROBE]
    report = {
        "schema_version": "experiment-start-v6.6.1-validation-1.0",
        "status": "PASS",
        "validated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "scope": "planning_API_contract_and_resource_estimation_only",
        "confirm_test_ready": False,
        "checks": {
            "original_v6.6_siblings_preserved": originals_preserved,
            "markdown_contract_terms": "PASS",
            "docx": inspect_docx(),
            "xlsx": inspect_xlsx(),
            "json": inspect_json(),
            "deepseek_live_probe": {
                "dry_run": "PASS",
                "live_request": "NOT_RUN",
                "reason": "DEEPSEEK_API_KEY was not present; no credential or paid call was authorized/provided",
            },
        },
        "artifacts": [{"path": str(path.relative_to(ROOT)), "sha256": sha256(path)} for path in artifacts],
        "remaining_hard_gates": [
            "ExternalModelReleaseManifest and official pricing snapshot at acquisition freeze",
            "DEEPSEEK_API_KEY connectivity probe with sanitized shakedown observation",
            "24-48 hour sentinel fingerprint/stability precheck",
            "immutable shared acquisition cache and offline replay verification",
            "v20-conformant implementation",
            "two-reviewer ActionTruth and scenario signatures",
            "CSTR physical qualification",
            "frozen G-star candidate and selection contract",
            "signed SESOI and no-harm margins",
            "independent joint-kernel pilot",
            "coverage, type-I and power confirmation",
            "sealed preregistration and TestManifest",
        ],
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(REPORT)


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
import re
import sys
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

import build_v16_delivery as base  # noqa: E402
from deliver_v20_v65 import citation_audit  # noqa: E402
from build_v201_docx_delivery import (  # noqa: E402
    fenced_code_to_boxes,
    insert_main_figure,
    structural_check,
)


JOBS = [
    ("00-论文-中文阅读稿v21-生命周期一致组合前治理与IF创新主线版.md", "v21：生命周期一致组合前治理"),
    ("实验设计与执行计划v6.8-机制先行与表示中立验证蓝图.md", "实验设计与执行计划 v6.8"),
    ("研判报告二十四-报告二十三裁决与v21-Information-Fusion创新性论证.md", "研判报告二十四"),
    ("实验执行记录v6.7-SHK-000-远程资格启动.md", "SHK-000 远程资格启动记录"),
]


def main() -> None:
    failures: list[str] = []
    conversion: dict[str, object] = {}
    with tempfile.TemporaryDirectory(prefix="v21-docx-") as tmp:
        tmpdir = Path(tmp)
        for filename, title in JOBS:
            source = ROOT / filename
            if not source.exists():
                raise FileNotFoundError(source)
            transformed, blocks = fenced_code_to_boxes(source.read_text(encoding="utf-8"))
            input_md = tmpdir / filename
            input_md.write_text(transformed, encoding="utf-8")
            target = source.with_suffix(".docx")
            conversion[target.name] = {
                **base.markdown_to_docx(input_md, target, title),
                "fenced_blocks_transformed": blocks,
            }

    main_docx = (ROOT / JOBS[0][0]).with_suffix(".docx")
    figure = ROOT / "figures" / "图1-安全工件到在线融合闭环-v20.png"
    figure_stats = insert_main_figure(main_docx, figure)

    files: dict[str, object] = {}
    for filename, _ in JOBS:
        md = ROOT / filename
        docx = md.with_suffix(".docx")
        check = structural_check(md, docx)
        files[docx.name] = check
        if check["zip_test"] is not None:
            failures.append(f"{docx.name}: zip test failed")
        if check["comments"] or check["tracked_insertions"] or check["tracked_deletions"]:
            failures.append(f"{docx.name}: comments/tracked changes remain")
        if check["fence_markers"]:
            failures.append(f"{docx.name}: fenced markers remain")

    main_text = (ROOT / JOBS[0][0]).read_text(encoding="utf-8")
    citation = citation_audit(main_text)
    abstract = re.search(r"## Abstract\s+(.*?)\s+Keywords:", main_text, re.S)
    words = re.findall(r"\b[A-Za-z][A-Za-z0-9–'’-]*\b", abstract.group(1) if abstract else "")
    keywords = re.search(r"(?m)^Keywords:\s*(.+)$", main_text)
    keyword_count = len([x for x in (keywords.group(1).split(";") if keywords else []) if x.strip()])
    if citation["reference_count"] > 50 or not citation["sequence_ok"]:
        failures.append("main reference list exceeds 50 or is not sequential")
    if citation["missing_entries"] or citation["uncited_references"]:
        failures.append("main has missing or uncited references")
    if len(words) > 250:
        failures.append("English abstract exceeds 250 words")
    if not 1 <= keyword_count <= 7:
        failures.append("English keyword count outside 1..7")
    if figure_stats["width_cm"] > 16.2:
        failures.append("main figure too wide")
    if not files[main_docx.name]["media"]:
        failures.append("main DOCX has no embedded media")

    required_main_tokens = [
        "evidence-specific exactly-once",
        "post-commit equivalence",
        "Registry-relative exclusion consistency",
        "claim_support=none",
        "Minimal Mechanism Experiment",
    ]
    doc = Document(main_docx)
    doc_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )
    for token in required_main_tokens:
        if token not in doc_text:
            failures.append(f"main DOCX missing token: {token}")

    report = {
        "status": "PASS" if not failures else "FAIL",
        "delivery_class": "READING_COPY_NOT_FINAL_OMML_SUBMISSION_SOURCE",
        "conversion": conversion,
        "files": files,
        "main_citation_audit": citation,
        "english_abstract_words": len(words),
        "english_keyword_count": keyword_count,
        "main_figure": {"path": str(figure), **figure_stats},
        "failures": failures,
    }
    out = ROOT / "v21-v6.8-DOCX结构与引用核验报告.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(out), "status": report["status"], "failures": failures}, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

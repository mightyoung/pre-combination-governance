from __future__ import annotations

import hashlib
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import build_v16_delivery as base  # noqa: E402
from deliver_v20_v65 import citation_audit, inspect_docx  # noqa: E402


JOBS = [
    ("00-论文-中文阅读稿v20.1-一致性修复与投稿主线收敛版.md", "v20.1：谱系依赖证据的组合前治理"),
    ("补充材料S1-v20.1-组合前治理证明与性质测试.md", "补充材料 S1 v20.1"),
    ("补充材料S2-v20.1-运行合同与观察映射规范.md", "补充材料 S2 v20.1"),
    ("实验设计与执行计划v6.7-精简确认性预注册蓝图.md", "实验设计与执行计划 v6.7"),
    ("研判报告二十二-v20与实验v6.6复审-Information-Fusion硬伤与v20.1-v6.7改进.md", "研判报告二十二"),
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def fenced_code_to_boxes(text: str) -> tuple[str, int]:
    """Turn fenced code into consecutive blockquote lines for the existing converter."""
    out: list[str] = []
    in_code = False
    block_count = 0
    for line in text.splitlines():
        if line.strip().startswith("```"):
            if not in_code:
                block_count += 1
                out.append("> CODE BLOCK")
            in_code = not in_code
            if not in_code:
                out.append("")
            continue
        if in_code:
            out.append("> " + (line if line else " "))
        else:
            out.append(line)
    if in_code:
        raise AssertionError("unclosed fenced code block")
    return "\n".join(out) + "\n", block_count


def insert_main_figure(docx_path: Path, image_path: Path) -> dict[str, float | int]:
    marker = "\u56fe1\uff1a\u5b89\u5168\u5de5\u4ef6\u5230\u5728\u7ebf\u6eaf\u6e90\u7ea6\u675f\u878d\u5408\u95ed\u73af"
    doc = Document(docx_path)
    matches = [p for p in doc.paragraphs if marker in p.text]
    if len(matches) != 1:
        raise AssertionError(f"expected one figure marker, got {len(matches)}")
    p = matches[0]
    for run in p.runs:
        run.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Cm(16.0))
    doc.save(docx_path)
    reopened = Document(docx_path)
    if len(reopened.inline_shapes) != 1:
        raise AssertionError(f"expected one inline figure, got {len(reopened.inline_shapes)}")
    shape = reopened.inline_shapes[0]
    return {
        "inline_shapes": len(reopened.inline_shapes),
        "width_cm": shape.width.cm,
        "height_cm": shape.height.cm,
    }


def structural_check(md_path: Path, docx_path: Path) -> dict[str, object]:
    doc = Document(docx_path)
    with zipfile.ZipFile(docx_path) as zf:
        bad = zf.testzip()
        names = set(zf.namelist())
        xml = zf.read("word/document.xml").decode("utf-8")
    return {
        **inspect_docx(docx_path),
        "source_md": md_path.name,
        "source_sha256": sha256(md_path),
        "docx_sha256": sha256(docx_path),
        "zip_test": bad,
        "reopen_ok": True,
        "comments": "word/comments.xml" in names,
        "tracked_insertions": xml.count("<w:ins"),
        "tracked_deletions": xml.count("<w:del"),
        "fence_markers": xml.count("```"),
        "omml_nodes": xml.count("<m:oMath"),
        "plain_math_classification": "READING_COPY_WITH_EDITABLE_PLAIN_MATH" if "<m:oMath" not in xml else "OMML_PRESENT",
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": sum(p.style.name.startswith("Heading") for p in doc.paragraphs),
    }


def main() -> None:
    for filename, _ in JOBS:
        if not (ROOT / filename).exists():
            raise FileNotFoundError(ROOT / filename)

    conversion: dict[str, object] = {}
    with tempfile.TemporaryDirectory(prefix="v201-docx-") as tmp:
        tmpdir = Path(tmp)
        for filename, short_title in JOBS:
            md_path = ROOT / filename
            input_md = md_path
            extra: dict[str, object] = {}
            if filename.startswith("补充材料S2-"):
                transformed, blocks = fenced_code_to_boxes(md_path.read_text(encoding="utf-8"))
                input_md = tmpdir / md_path.name
                input_md.write_text(transformed, encoding="utf-8")
                extra["fenced_blocks_transformed"] = blocks
            docx_path = md_path.with_suffix(".docx")
            conversion[docx_path.name] = {
                **base.markdown_to_docx(input_md, docx_path, short_title),
                **extra,
            }

    main_docx = (ROOT / JOBS[0][0]).with_suffix(".docx")
    figure = ROOT / "figures" / "图1-安全工件到在线融合闭环-v20.png"
    figure_stats = insert_main_figure(main_docx, figure)

    files: dict[str, object] = {}
    failures: list[str] = []
    for filename, _ in JOBS:
        md_path = ROOT / filename
        docx_path = md_path.with_suffix(".docx")
        check = structural_check(md_path, docx_path)
        files[docx_path.name] = check
        if check["zip_test"] is not None:
            failures.append(f"{docx_path.name}: zip test failed")
        if check["comments"] or check["tracked_insertions"] or check["tracked_deletions"]:
            failures.append(f"{docx_path.name}: comments or tracked changes remain")
        if check["fence_markers"]:
            failures.append(f"{docx_path.name}: fenced code marker remains")

    s2_doc = Document((ROOT / JOBS[2][0]).with_suffix(".docx"))
    s2_text = "\n".join(p.text for p in s2_doc.paragraphs) + "\n" + "\n".join(
        cell.text for table in s2_doc.tables for row in table.rows for cell in row.cells
    )
    for token in ["attempt_status", "SetwiseEligibilityCertificate", "ABSTAIN_MAPPING"]:
        if token not in s2_text:
            failures.append(f"S2 code/contract token missing after conversion: {token}")

    main_text = (ROOT / JOBS[0][0]).read_text(encoding="utf-8")
    citation = citation_audit(main_text)
    english_match = re.search(r"## Abstract\s+(.*?)\s+Keywords:", main_text, re.S)
    english_words = re.findall(r"\b[A-Za-z][A-Za-z0-9–'’-]*\b", english_match.group(1) if english_match else "")
    keywords_match = re.search(r"(?m)^Keywords:\s*(.+)$", main_text)
    keyword_count = len([x for x in (keywords_match.group(1).split(";") if keywords_match else []) if x.strip()])
    if citation["reference_count"] != 50 or not citation["sequence_ok"]:
        failures.append("main citation sequence is not exactly 1..50")
    if citation["missing_entries"] or citation["uncited_references"]:
        failures.append("main has missing or uncited references")
    if len(english_words) > 250:
        failures.append("English abstract exceeds 250 words")
    if not 1 <= keyword_count <= 7:
        failures.append("English keywords outside 1..7")
    if figure_stats["width_cm"] > 16.2:
        failures.append("main figure exceeds available page width")
    if not files[main_docx.name]["media"]:
        failures.append("main DOCX has no embedded media")

    report = {
        "status": "PASS" if not failures else "FAIL",
        "delivery_class": "READING_COPY_NOT_FINAL_OMML_SUBMISSION_SOURCE",
        "conversion": conversion,
        "files": files,
        "main_citation_audit": citation,
        "english_abstract_words": len(english_words),
        "english_keyword_count": keyword_count,
        "main_figure": {"path": str(figure), "sha256": sha256(figure), **figure_stats},
        "failures": failures,
    }
    out = ROOT / "v20.1-v6.7-DOCX结构核验报告.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(out), "status": report["status"], "failures": failures}, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

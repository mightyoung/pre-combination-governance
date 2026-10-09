#!/usr/bin/env python3
"""Apply the conservative T4-aware capacity envelope to the v6.6 workbook."""

from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "tools" / "build_experiment_resource_workbook_v66.py"
OUT = ROOT / "实验资源与样本量预算v6.6.xlsx"


def run_base() -> None:
    spec = importlib.util.spec_from_file_location("v66_base", BASE)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load base workbook builder")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.main()


def main() -> None:
    run_base()
    wb = load_workbook(OUT)
    readme = wb["说明"]
    readme["B6"] = "32 逻辑线程、128 GB RAM、2 TB 可用 NVMe、1×NVIDIA T4"
    readme["C6"] = "约 169 人日"
    readme["D6"] = "约 120 T4 GPU-h；约 1.2k CPU-h 中央估计。"
    readme["B7"] = "功效初始工作域约2,500–10,000个可分析bundle；资源按14,400预留"
    readme["C7"] = "推荐容量点 n_H/n_G/n_B=9,600/2,400/2,400"
    readme["D7"] = "正式 n 只允许在独立 pilot 与 MC confirm 后冻结；n_max_generate 建议24,000。"

    inputs = wb["假设输入"]
    replacements = {
        "hazard_confirm_bundles": (4800, 9600, 14400),
        "benign_confirm_bundles": (1200, 2400, 3600),
        "budget_confirm_bundles": (1200, 2400, 3600),
        "hazard_methods": (4, 5, 6),
        "benign_methods": (2, 3, 4),
        "budget_methods": (2, 5, 6),
        "variants": (1, 3, 5),
        "algorithm_seeds": (1, 2, 3),
        "pretest_method_runs": (12000, 144000, 864000),
        "external_method_runs": (3600, 90000, 720000),
        "power_sim_cpu_hours": (50, 250, 1000),
        "retained_mb_per_method_run": (0.1, 0.3, 0.5),
        "gpu_hours": (0, 120, 300),
    }
    for row in range(4, inputs.max_row + 1):
        name = inputs.cell(row, 1).value
        if name in replacements:
            for col, value in zip(range(2, 5), replacements[name]):
                inputs.cell(row, col, value)

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

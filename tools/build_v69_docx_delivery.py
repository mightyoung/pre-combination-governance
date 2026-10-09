from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

import build_v16_delivery as base  # noqa: E402
from build_v201_docx_delivery import fenced_code_to_boxes, structural_check  # noqa: E402


SOURCE = ROOT / "实验设计与执行计划v6.9-指定执行版-机制验证到确认性试验.md"
TARGET = SOURCE.with_suffix(".docx")
REPORT = ROOT / "v6.9-指定实验计划交付与结构核验.json"


def document_text(path: Path) -> str:
    doc = Document(path)
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    return paragraphs + "\n" + cells


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    with tempfile.TemporaryDirectory(prefix="v69-docx-") as tmp:
        transformed, blocks = fenced_code_to_boxes(SOURCE.read_text(encoding="utf-8"))
        converted_source = Path(tmp) / SOURCE.name
        converted_source.write_text(transformed, encoding="utf-8")
        conversion = {
            **base.markdown_to_docx(
                converted_source,
                TARGET,
                "实验设计与执行计划 v6.9（指定执行版）",
            ),
            "fenced_blocks_transformed": blocks,
        }

    check = structural_check(SOURCE, TARGET)
    text = document_text(TARGET)
    required_tokens = [
        "physical_seed_base_bundle",
        "MME-SCREEN",
        "96 bundles",
        "MME-VERIFY",
        "17,472",
        "PENDING_DOMAIN_SIGNOFF",
        "10,000",
        "D_test",
        "REMOTE-RUN-BLOCKED",
        "joint power",
    ]

    failures: list[str] = []
    if check["zip_test"] is not None:
        failures.append("DOCX ZIP integrity test failed")
    if check["comments"] or check["tracked_insertions"] or check["tracked_deletions"]:
        failures.append("comments or tracked changes remain")
    if check["fence_markers"]:
        failures.append("fenced-code markers remain")
    for token in required_tokens:
        if token not in text:
            failures.append(f"DOCX missing required token: {token}")

    report = {
        "status": "PASS" if not failures else "FAIL",
        "delivery_class": "EXECUTABLE_EXPERIMENT_PLAN_READING_COPY",
        "source": str(SOURCE),
        "target": str(TARGET),
        "conversion": conversion,
        "structure": check,
        "required_tokens": required_tokens,
        "failures": failures,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "docx": str(TARGET),
                "report": str(REPORT),
                "status": report["status"],
                "failures": failures,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

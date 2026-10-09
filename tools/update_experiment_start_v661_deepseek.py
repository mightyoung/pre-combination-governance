#!/usr/bin/env python3
"""Create the v6.6.1 experiment-start siblings with DeepSeek V4 Flash API planning.

The v6.6 artifacts are preserved.  This script updates the Markdown, checklist,
and resource workbook into new v6.6.1 siblings.
"""

from __future__ import annotations

import json
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "实验启动计划v6.6-双机执行与资源预算.md"
OUT_MD = ROOT / "实验启动计划v6.6.1-DeepSeek在线API与双机资源预算.md"
SRC_XLSX = ROOT / "实验资源与样本量预算v6.6.xlsx"
OUT_XLSX = ROOT / "实验资源与样本量预算v6.6.1.xlsx"
SRC_JSON = ROOT / "实验启动清单v6.6.json"
OUT_JSON = ROOT / "实验启动清单v6.6.1.json"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
thin = Side(style="thin", color="B7B7B7")


API_SECTION = r"""
### 8.2 DeepSeek V4 Flash 在线 API 的冻结合同

DeepSeek 只允许作为**公共、冻结的外部随机 SourceAdapter**，不允许成为 A 的私有能力。正确数据流是：

`raw operational observation → one frozen DeepSeek acquisition → immutable structured-evidence cache → identical payload to A/G*/C/ablations/M0–M4`。

错误且会使主比较失效的数据流是：A 在线调用 DeepSeek，而 G* 使用本地模型、旧缓存或另一轮独立抽样。那样测到的是治理机制、模型能力、供应商漂移和 API 运气的混合效应，不能归因于组合前治理。

| 合同字段 | v6.6.1 规划值 | 冻结要求 |
| --- | --- | --- |
| 接口/模型 | OpenAI 格式 `https://api.deepseek.com`；`deepseek-v4-flash` | 禁止使用即将弃用的 `deepseek-chat` / `deepseek-reasoner` 别名 |
| 模式 | `thinking.type=disabled` | 官方默认是 enabled；必须显式关闭 |
| 采样参数 | `temperature=0`, `top_p=1`, `stream=false`, `tools=[]` | 无公开 seed；这些参数不等于逐 token 确定性 |
| 结构化输出 | `response_format={"type":"json_object"}`；本地 JSON Schema + 语义 validator | prompt 必须含 json 与输出示例；合法 JSON 仍不等于业务合法 |
| 输出帽 | `max_tokens=2048` | vertical slice 后按 P99 输出长度重签；不得在 test 中临时放宽 |
| 并发 | 推荐 32；微基准范围 8–64 | 官方账户上限 2,500，但不得为追求吞吐制造时间块/429 混杂 |
| deadline | connect 10 s；单请求总 deadline 180 s | 迟到响应不得回填；采集延迟与方法 deadline 分开 |
| 重试 | 仅在未收到响应体的 429/500/503 上，完全相同请求最多重试 2 次，冻结退避 2 s/8 s | 200 空内容、非法 schema、length、content_filter、资源不足均不“刷到成功” |
| 密钥 | 环境变量 `DEEPSEEK_API_KEY` | 禁止写入代码、manifest、日志、DOCX/XLSX 或容器层 |
| 模型版本证据 | 保存 response `model`、`system_fingerprint`、response/request ID、created、全部 headers 摘要 | fingerprint 变化立即停止当前 acquisition version；不得只补跑受影响样本 |
| token/缓存证据 | 保存 prompt/completion tokens、`prompt_cache_hit_tokens`、`prompt_cache_miss_tokens` | 供应商缓存只影响成本/延迟，不是科学工件缓存，也不保证输出确定性 |

确认性 acquisition 必须在短时间窗内按 topology/path/severity/namespace 完整区组随机交错，不能先跑严重事故再跑良性场景。正式方法运行前封存 `ExternalModelReleaseManifest`、canonical request、prompt/schema SHA-256、SDK 版本、响应原文、规范化 JSON、失败状态和内容寻址 cache key；之后关闭网络，所有方法只读同一缓存。论文只能声称“下游工件可重放，结论条件于指定采集日期、fingerprint 集合与冻结缓存”，不能声称未来重新调用 API 会得到相同输出。

信息隔离是硬门：请求不得包含 TruthManifest、科学层/严重度真值、真实 CCF/屏障成败、未来状态、正确动作、endpoint、方法臂身份、scorer 反馈或既往 test 结果。每次请求无跨 episode 会话历史，只使用当时因果可获得的 operational observation 和公开 OperationalSafetyArtifact；`user_id` 只能使用不含隐私和真值的 namespace 伪标识。

独立统计单位仍是 physical-seed base bundle，不是 prompt、response、API retry 或 K 次重复。pilot 可在预签子集上用 K=3 估计语义映射不稳定性；confirm 主分析默认每个冻结 emitter/claim group 接受首个合格响应，任何 K 与 reducer 必须在 G6 前固定，且所有方法共享同一组 K 输出。API failure、schema failure、fingerprint/time block 与 K 内相关必须进入 pilot kernel、覆盖和功效敏感性。

### 8.3 API 调用量、token 与费用预算

DeepSeek 当前人民币单价为：输入缓存命中 0.02 元/百万 token、输入缓存未命中 1 元/百万 token、输出 2 元/百万 token。缓存默认开启但为 best-effort，因此预算同时给出预计命中成本和完全未命中成本。下面的“调用/bundle”表示五证据流经冻结 emitter 分组后的调用等价数，不是方法数；vertical slice 后必须替换。

| 项目 | 低档 | 推荐 | 高档 |
| --- | ---: | ---: | ---: |
| API-eligible unique bundles | 10,000 | 24,000 | 48,000 |
| emitter 调用/bundle | 1 | 2 | 4 |
| 基础调用 | 10,000 | 48,000 | 192,000 |
| 额外计费尝试比例 | 5% | 10% | 20% |
| 计费尝试 | 10,500 | 52,800 | 230,400 |
| 平均输入/输出 token | 2,000 / 300 | 3,000 / 500 | 5,000 / 1,000 |
| 规划缓存命中率 | 50% | 70% | 80% |
| 预计费用（人民币） | 约 17.01 元 | 约 102.54 元 | 约 709.63 元 |
| 完全无缓存费用 | 约 27.30 元 | 约 211.20 元 | 约 1,612.80 元 |
| 建议授权/充值帽 | 100 元 | 1,000 元 | 5,000 元 |

推荐档授权帽明显高于中央估计，用于覆盖价格变化、缓存失效、长输出、失败尝试和扩大支持性分析，并不意味着应花完。正式采集前重新保存官方价格页快照；若价格、模型默认模式或 API schema 变化，必须升版 `ExternalModelReleaseManifest`。稳定 system/schema/few-shot 前缀应置前，动态 observation 置后，以提升供应商缓存命中；但不得为了省钱改变科学随机化、隔离或失败合同。
""".strip()


def replace_once(text: str, old: str, new: str) -> str:
    if old not in text:
        raise RuntimeError(f"Expected text not found: {old[:100]!r}")
    return text.replace(old, new, 1)


def build_markdown() -> None:
    text = SRC_MD.read_text(encoding="utf-8")
    text = replace_once(text, "# 实验启动计划 v6.6：双机执行、功效路线与资源预算",
                        "# 实验启动计划 v6.6.1：DeepSeek V4 Flash 在线 API、双机执行与资源预算")
    text = replace_once(text, "- Version Label: experiment_start_v6.6",
                        "- Version Label: experiment_start_v6.6.1")
    text = replace_once(text, "本文件是 v6.5 的执行与资源工作层",
                        "本文件是 v6.5 的 v6.6.1 执行与资源工作层")
    text = replace_once(
        text,
        "- T4 服务器：已知配备 NVIDIA T4；CPU、RAM、NVMe、系统与驱动待实测。承担容器化批量运行、IFNet/AE/LSTM 等神经基线和确认性作业。",
        "- T4 服务器：已知配备 NVIDIA T4；CPU、RAM、NVMe、系统与驱动待实测。承担容器化批量运行、IFNet/AE/LSTM 等非大模型神经基线和确认性作业；不再承担大模型推理。\n- 在线大模型：统一使用 DeepSeek 官方 `deepseek-v4-flash` API，先采集并冻结公共语义证据缓存，再供全部方法只读复用；API 不进入方法私有服务合同。",
    )
    text = replace_once(text, "| T4 GPU-hours | 0–20 | 约 100–300 | 300–1,000 |",
                        "| T4 GPU-hours | 0–8 | 约 40–120（中央 80） | 150–400 |\n| DeepSeek API 费用 | 约 17 元；授权帽 100 元 | 约 103 元；授权帽 1,000 元 | 约 710 元；授权帽 5,000 元 |")
    text = replace_once(text, "## 8. T4 服务器的正确用途",
                        "## 8. 计算分工：T4 与 DeepSeek 在线 API\n\n### 8.1 T4 服务器的正确用途")
    text = replace_once(text, "- Observation adapter 或表示层 C 的冻结批量推理；",
                        "- 表示层 C 以及不依赖在线大模型的冻结批量推理；")
    text = replace_once(
        text,
        "旧 README 中的本地 Qwen2.5-VL-72B 路线不适合单卡 T4 16 GB，也不属于当前 v20 的确认性主线。确认阶段不得临时改量化、micro-batch 或精度；若需要这些设置，必须在 G4 前冻结并通过动作输出等价测试。",
        "旧 README 中的本地 Qwen2.5-VL-72B 路线不适合单卡 T4 16 GB，也不属于当前 v20 的确认性主线。在线 DeepSeek 已取代本地大模型 emitter；T4 的量化、micro-batch 或精度仅适用于非大模型神经基线，并必须在 G4 前冻结并通过动作输出等价测试。\n\n" + API_SECTION,
    )
    text = replace_once(
        text,
        "8. 用 100–300 bundle 和 1,000 episode 记录 generator/replay P50/P95、RAM、磁盘/episode、T4吞吐与失败码；据此更新资源工作簿。",
        "8. 用 100–300 bundle 和 1,000 episode 记录 generator/replay P50/P95、RAM、磁盘/episode、T4吞吐与失败码；另以脱敏 shakedown observation 验证 DeepSeek canonical request、JSON/schema failure、token、cache hit/miss、fingerprint、并发 8/32/64 和采集延迟；据此更新资源工作簿。",
    )
    text = replace_once(
        text,
        "从第一条实验记录开始保存：数据 hash、代码 commit、release DAG、容器 digest、seed namespace、配置、CPU/GPU、driver/CUDA、deterministic flags、方法 token、失败码和输出 hash。",
        "从第一条实验记录开始保存：数据 hash、代码 commit、release DAG、容器 digest、seed namespace、配置、CPU/GPU、driver/CUDA、deterministic flags、方法 token、失败码和输出 hash。对 DeepSeek 额外保存 ExternalModelReleaseManifest、prompt/schema/request hash、model、system_fingerprint、response/request ID、created、finish_reason、HTTP/重试、完整原始响应、规范化 JSON、token/cache 统计、延迟与不可覆盖 cache key。",
    )
    text = replace_once(
        text,
        "8. Rieth et al. Additional Tennessee Eastman Process Simulation Data. https://doi.org/10.7910/DVN/6C3JR1",
        "8. Rieth et al. Additional Tennessee Eastman Process Simulation Data. https://doi.org/10.7910/DVN/6C3JR1\n9. DeepSeek. Models & Pricing. https://api-docs.deepseek.com/zh-cn/quick_start/pricing/\n10. DeepSeek. Thinking Mode. https://api-docs.deepseek.com/zh-cn/guides/thinking_mode/\n11. DeepSeek. JSON Output. https://api-docs.deepseek.com/zh-cn/guides/json_mode/\n12. DeepSeek. Context Caching. https://api-docs.deepseek.com/zh-cn/guides/kv_cache/\n13. DeepSeek. Rate Limit & Isolation. https://api-docs.deepseek.com/quick_start/rate_limit/\n14. DeepSeek. Change Log. https://api-docs.deepseek.com/updates/",
    )
    OUT_MD.write_text(text, encoding="utf-8")


def set_title(ws, text: str) -> None:
    if ws.merged_cells.ranges:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row == 1:
                ws.unmerge_cells(str(merged))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    c = ws.cell(1, 1, text)
    c.font = Font(name="Arial", size=16, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def style_api_sheet(ws) -> None:
    set_title(ws, "DeepSeek V4 Flash API 调用量、token 与费用预算")
    headers = ["参数/指标", "低档", "推荐", "高档", "解释/冻结要求"]
    for col, value in enumerate(headers, 1):
        c = ws.cell(3, col, value)
        c.font = Font(name="Arial", bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=5):
        for c in row:
            c.font = copy(c.font)
            c.font = Font(name="Arial", size=10, color=c.font.color or "000000")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, width in {1: 38, 2: 16, 3: 16, 4: 16, 5: 68}.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def add_api_budget_sheet(wb) -> None:
    if "DeepSeek API预算" in wb.sheetnames:
        del wb["DeepSeek API预算"]
    ws = wb.create_sheet("DeepSeek API预算", 3)
    inputs = [
        ("api_eligible_unique_bundles", [10000, 24000, 48000], "train/select/pilot/test/external 中需要语义 emitter 的唯一 bundle；不是统计 n"),
        ("emitter_calls_per_bundle", [1, 2, 4], "五证据流经冻结 emitter 分组后的调用等价数；不得乘 method/variant/seed"),
        ("avg_input_tokens_per_attempt", [2000, 3000, 5000], "由 vertical slice 的 usage P50/P95 替换"),
        ("avg_output_tokens_per_attempt", [300, 500, 1000], "非思考 JSON emitter；max_tokens 规划值 2048"),
        ("provider_cache_hit_rate", [0.50, 0.70, 0.80], "best-effort 成本假设；不得视为科学复现缓存"),
        ("extra_billable_attempt_ratio", [0.05, 0.10, 0.20], "相对基础调用的额外计费尝试；含冻结的瞬态重试与支持性采集"),
        ("avg_latency_seconds", [3, 6, 12], "由 shakedown 实测替换；不计入下游方法 deadline"),
        ("requested_concurrency", [8, 32, 64], "官方账户上限 2500；实验主动限流避免时间块混杂"),
        ("concurrency_efficiency", [0.70, 0.70, 0.65], "排队、尾延迟与失败损失"),
        ("input_cache_hit_RMB_per_M", [0.02, 0.02, 0.02], "DeepSeek 官方价格，核对日 2026-07-14"),
        ("input_cache_miss_RMB_per_M", [1.00, 1.00, 1.00], "DeepSeek 官方价格，核对日 2026-07-14"),
        ("output_RMB_per_M", [2.00, 2.00, 2.00], "DeepSeek 官方价格，核对日 2026-07-14"),
    ]
    row = 4
    for name, values, note in inputs:
        ws.cell(row, 1, name)
        ws.cell(row, 5, note)
        for col, value in enumerate(values, 2):
            c = ws.cell(row, col, value)
            c.font = Font(name="Arial", size=10, color=INPUT_BLUE)
            c.fill = PatternFill("solid", fgColor=YELLOW)
            c.comment = Comment("Planning input; replace from frozen vertical-slice measurements or an updated official price snapshot.", "Codex")
        row += 1

    formula_rows = [
        ("base_API_calls", "={c}4*{c}5", "唯一 bundle × emitter 调用；不乘方法、variant 或 seed"),
        ("billable_attempts", "={c}16*(1+{c}9)", "基础调用加规划额外计费尝试"),
        ("input_tokens_M", "={c}17*{c}6/1000000", "输入 token 百万量"),
        ("output_tokens_M", "={c}17*{c}7/1000000", "输出 token 百万量"),
        ("cache_hit_input_tokens_M", "={c}18*{c}8", "仅为供应商计费/延迟统计"),
        ("cache_miss_input_tokens_M", "={c}18*(1-{c}8)", "完全无缓存敏感性另列"),
        ("estimated_cost_RMB", "={c}20*{c}13+{c}21*{c}14+{c}19*{c}15", "按当前官方人民币价格"),
        ("worst_case_no_cache_cost_RMB", "={c}18*{c}14+{c}19*{c}15", "provider cache 0% 命中"),
        ("ideal_acquisition_wall_hours", "={c}17*{c}10/(3600*{c}11*{c}12)", "不含阶段依赖和 fingerprint 变化导致的整版作废"),
        ("authorization_cap_RMB", None, "建议授权/充值硬帽；价格或模型合同变化时重新签署"),
    ]
    for name, formula, note in formula_rows:
        ws.cell(row, 1, name)
        ws.cell(row, 5, note)
        for col in range(2, 5):
            letter = get_column_letter(col)
            if formula:
                ws.cell(row, col, formula.format(c=letter))
        row += 1
    for col, value in enumerate([100, 1000, 5000], 2):
        ws.cell(25, col, value)
        ws.cell(25, col).font = Font(name="Arial", size=10, color=INPUT_BLUE)
        ws.cell(25, col).fill = PatternFill("solid", fgColor=YELLOW)
    ws.cell(26, 1, "provider_account_concurrency_limit")
    ws.cell(26, 5, "DeepSeek 官方账户级并发；所有 API key 合并计数；本实验不应接近该上限")
    for col in range(2, 5):
        ws.cell(26, col, 2500)
    ws.cell(28, 1, "官方资料")
    ws.cell(28, 2, "模型与价格")
    ws.cell(28, 2).hyperlink = "https://api-docs.deepseek.com/zh-cn/quick_start/pricing/"
    ws.cell(29, 2, "思考模式")
    ws.cell(29, 2).hyperlink = "https://api-docs.deepseek.com/zh-cn/guides/thinking_mode/"
    ws.cell(30, 2, "JSON Output")
    ws.cell(30, 2).hyperlink = "https://api-docs.deepseek.com/zh-cn/guides/json_mode/"
    ws.cell(31, 2, "上下文缓存")
    ws.cell(31, 2).hyperlink = "https://api-docs.deepseek.com/zh-cn/guides/kv_cache/"
    ws.cell(32, 2, "并发限制")
    ws.cell(32, 2).hyperlink = "https://api-docs.deepseek.com/quick_start/rate_limit/"
    style_api_sheet(ws)
    for row_idx in range(4, 27):
        for col in range(2, 5):
            ws.cell(row_idx, col).number_format = "#,##0.00"
    for row_idx in [8, 9, 12]:
        for col in range(2, 5):
            ws.cell(row_idx, col).number_format = "0.0%"
    for col in range(2, 5):
        ws.cell(22, col).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        ws.cell(23, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)


def update_workbook() -> None:
    wb = load_workbook(SRC_XLSX)
    info = wb["说明"]
    info["A1"] = "实验资源与样本量预算 v6.6.1"
    info["B4"] = "PLANNING-ONLY；已加入 DeepSeek V4 Flash 在线 API；不是正式功效冻结或 TestManifest"
    info["B5"] = "DeepSeek 作为统一预采集并冻结的公共 SourceAdapter；T4 仅承担非大模型神经基线。"
    info["C5"] = "API 不能按 method/variant/seed 分别调用；方法运行前关闭网络并共享同一证据缓存。"
    info["D5"] = "先冻结 ExternalModelReleaseManifest、fingerprint、prompt/schema hash 与失败合同。"
    info["D6"] = "中央约 80 T4 GPU-h；DeepSeek 预计约 103 元，建议 API 授权帽 1,000 元。"

    inputs = wb["假设输入"]
    for row in range(4, inputs.max_row + 1):
        if inputs.cell(row, 1).value == "gpu_hours":
            for col, value in zip(range(2, 5), [0, 80, 400]):
                inputs.cell(row, col, value)
            inputs.cell(row, 5, "仅 IFNet/AE/LSTM/evidential-DL/表示层 C 等非大模型神经基线；在线 DeepSeek 不占 T4")

    total = wb["总资源"]
    new_row = total.max_row + 1
    total.cell(new_row, 1, "DeepSeek_API_authorization_cap_RMB")
    total.cell(new_row, 5, "授权/充值硬帽，不是预计消费；预计成本见 DeepSeek API预算")
    for col in range(2, 5):
        letter = get_column_letter(col)
        total.cell(new_row, col, f"='DeepSeek API预算'!{letter}25")
        total.cell(new_row, col).number_format = "#,##0.00"
        total.cell(new_row, col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
        total.cell(new_row, col).alignment = Alignment(vertical="top", wrap_text=True)
    for col in [1, 5]:
        total.cell(new_row, col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
        total.cell(new_row, col).alignment = Alignment(vertical="top", wrap_text=True)

    add_api_budget_sheet(wb)

    sources = wb["研究依据"]
    append_rows = [
        ("DeepSeek 模型/价格", "官方文档", "deepseek-v4-flash；价格、上下文、输出、别名弃用", "https://api-docs.deepseek.com/zh-cn/quick_start/pricing/", "2026-07-14 核对；正式采集前重新快照"),
        ("DeepSeek 思考模式", "官方文档", "默认 thinking enabled；显式 disabled", "https://api-docs.deepseek.com/zh-cn/guides/thinking_mode/", "非思考 JSON emitter 合同"),
        ("DeepSeek JSON Output", "官方文档", "JSON mode、空 content 与截断风险", "https://api-docs.deepseek.com/zh-cn/guides/json_mode/", "必须本地 schema/语义校验"),
        ("DeepSeek Context Cache", "官方文档", "best-effort 输入前缀缓存与 usage 字段", "https://api-docs.deepseek.com/zh-cn/guides/kv_cache/", "只影响成本/延迟，不保证确定性"),
        ("DeepSeek Rate Limit", "官方文档", "Flash 账户级并发 2500", "https://api-docs.deepseek.com/quick_start/rate_limit/", "实验主动限流 8–64"),
        ("DeepSeek Change Log", "官方文档", "V4 Flash 与旧别名弃用", "https://api-docs.deepseek.com/updates/", "禁止旧别名"),
    ]
    for values in append_rows:
        row = sources.max_row + 1
        for col, value in enumerate(values, 1):
            sources.cell(row, col, value)
            sources.cell(row, col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
            sources.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        sources.cell(row, 4).hyperlink = values[3]
        sources.cell(row, 4).style = "Hyperlink"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT_XLSX)


def update_checklist() -> None:
    data = json.loads(SRC_JSON.read_text(encoding="utf-8"))
    data["schema_version"] = "experiment-start-v6.6.1-1.0"
    data["release_binding"]["execution_plan"] = "v6.6.1"
    data["declared_server"]["role"] = "non_LLM_neural_baselines_and_batch_confirmation_only"
    data["recommended_server_threshold"]["gpu"] = "1x NVIDIA T4 is sufficient for planned non-LLM neural baselines"
    data["recommended_resource_envelope"]["t4_gpu_hours"] = [40, 120]
    data["recommended_resource_envelope"]["t4_gpu_hours_central"] = 80
    data["recommended_resource_envelope"]["deepseek_api_expected_cost_rmb"] = 102.54
    data["recommended_resource_envelope"]["deepseek_api_no_cache_cost_rmb"] = 211.20
    data["recommended_resource_envelope"]["deepseek_api_authorization_cap_rmb"] = 1000
    data["online_llm_api"] = {
        "role": "shared_frozen_external_random_SourceAdapter",
        "model": "deepseek-v4-flash",
        "provider_release_status": "hosted_preview_without_public_date_snapshot_id",
        "openai_base_url": "https://api.deepseek.com",
        "anthropic_base_url": "https://api.deepseek.com/anthropic",
        "pricing_checked_date": str(date(2026, 7, 14)),
        "pricing_rmb_per_million_tokens": {
            "input_cache_hit": 0.02,
            "input_cache_miss": 1.0,
            "output": 2.0
        },
        "provider_limits": {
            "context_tokens": 1000000,
            "max_output_tokens": 384000,
            "account_concurrency": 2500
        },
        "frozen_request_contract": {
            "thinking": {"type": "disabled"},
            "temperature": 0,
            "top_p": 1,
            "stream": False,
            "tools": [],
            "response_format": {"type": "json_object"},
            "max_tokens": 2048,
            "requested_concurrency": 32,
            "connect_timeout_seconds": 10,
            "total_deadline_seconds": 180,
            "api_key_environment_variable": "DEEPSEEK_API_KEY"
        },
        "reproducibility_limits": [
            "no_public_seed_parameter",
            "no_public_date_pinned_snapshot_id",
            "temperature_zero_does_not_guarantee_identical_output",
            "provider_context_cache_does_not_cache_output",
            "claims_are_conditional_on_acquisition_window_fingerprints_and_frozen_response_cache"
        ],
        "mandatory_record_fields": [
            "canonical_request_sha256", "prompt_sha256", "schema_sha256", "model",
            "system_fingerprint", "response_id", "request_id", "created", "http_status",
            "finish_reason", "raw_response", "normalized_json", "schema_validation",
            "prompt_tokens", "completion_tokens", "prompt_cache_hit_tokens",
            "prompt_cache_miss_tokens", "latency_ms", "attempt_number", "cache_key"
        ],
        "retry_failure_contract": {
            "retry_only": "429_500_503_before_any_response_body",
            "maximum_retries": 2,
            "fixed_backoff_seconds": [2, 8],
            "http_200_empty_invalid_schema_length_content_filter": "MAPPING_UNKNOWN_NO_RETRY",
            "insufficient_system_resource_or_partial_body": "EXTERNAL_SERVICE_FAILURE_NO_RETRY",
            "http_400_422": "STOP_ACQUISITION_VERSION_CONFIGURATION_ERROR",
            "http_401_402": "STOP_BATCH_OWNER_OR_CREDIT_ERROR",
            "deadline_exceeded": "FAIL_LATE_RESPONSE_MUST_NOT_BACKFILL"
        },
        "fairness_and_isolation": {
            "preacquire_before_method_runs": True,
            "immutable_shared_cache_for_all_methods": True,
            "network_disabled_during_confirm_method_runs": True,
            "forbidden_request_content": [
                "TruthManifest", "scientific_stratum_or_severity_truth", "true_CCF_or_barrier_outcome",
                "future_state", "correct_action", "endpoint_or_scorer_feedback", "method_arm_identity"
            ],
            "independent_unit": "physical_seed_base_bundle_not_API_call_or_retry"
        },
        "recommended_budget": {
            "api_eligible_unique_bundles": 24000,
            "emitter_calls_per_bundle": 2,
            "base_calls": 48000,
            "extra_billable_attempt_ratio": 0.10,
            "billable_attempts": 52800,
            "average_input_tokens": 3000,
            "average_output_tokens": 500,
            "provider_cache_hit_rate_assumption": 0.70,
            "expected_cost_rmb": 102.54,
            "complete_cache_miss_cost_rmb": 211.20,
            "authorization_cap_rmb": 1000
        },
        "official_sources": [
            "https://api-docs.deepseek.com/zh-cn/quick_start/pricing/",
            "https://api-docs.deepseek.com/zh-cn/guides/thinking_mode/",
            "https://api-docs.deepseek.com/zh-cn/guides/json_mode/",
            "https://api-docs.deepseek.com/zh-cn/guides/kv_cache/",
            "https://api-docs.deepseek.com/quick_start/rate_limit/",
            "https://api-docs.deepseek.com/updates/"
        ]
    }
    data["gates"].insert(1, {
        "gate": "G0-API",
        "status": "BLOCKED",
        "output": "ExternalModelReleaseManifest, sentinel stability check, canonical request and immutable acquisition cache contract"
    })
    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    build_markdown()
    update_workbook()
    update_checklist()
    print(OUT_MD)
    print(OUT_XLSX)
    print(OUT_JSON)


if __name__ == "__main__":
    main()

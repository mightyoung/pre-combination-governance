from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "含能材料制造多源风险证据的生命周期一致组合前治理方法-演示稿.md"
OUTPUT = ROOT / "含能材料制造多源风险证据的生命周期一致组合前治理方法-兵工学报格式-演示稿.docx"


def set_east_asia(run, font_name: str) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_margins(cell, top=50, start=70, bottom=50, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="000000", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))
    run.font.size = Pt(8)


def configure_section(section, columns: int = 1) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.6)
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "566" if columns == 2 else "720")


def clean_inline(text: str) -> str:
    text = text.strip()
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = text.replace("`", "")
    return text


def add_text_run(paragraph, text: str, size=9, bold=False, italic=False, font="宋体", color=None):
    run = paragraph.add_run(text)
    set_east_asia(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_body_paragraph(doc, text: str, *, first_indent=True, italic=False, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(13.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_indent and not center:
        pf.first_line_indent = Pt(18)
    add_text_run(p, clean_inline(text), size=9, italic=italic)
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(14)
    add_text_run(p, text, size=10.5 if level == 1 else 9, bold=True, font="黑体")
    return p


def add_markdown_table(doc, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table, size="4")
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_text_run(p, clean_inline(value), size=7.5, bold=(r_idx == 0), font="宋体")
            if r_idx == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7ECF2")
                cell._tc.get_or_add_tcPr().append(shd)
        if r_idx == 0:
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def parse_table(lines: list[str], start: int):
    raw = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw.append(lines[i].strip())
        i += 1
    rows = []
    for idx, line in enumerate(raw):
        cells = [x.strip() for x in line.strip("|").split("|")]
        if idx == 1 and all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
            continue
        rows.append(cells)
    return rows, i


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_section(doc.sections[0], columns=1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(9)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(13.5)

    i = 0
    in_cn_abstract = False
    in_en_abstract = False
    body_started = False
    front_index = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(26)
            add_text_run(p, line[2:], size=22, bold=True, font="宋体")
            i += 1
            continue

        if line == "## 摘要":
            in_cn_abstract = True
            in_en_abstract = False
            i += 1
            continue
        if line == "## Abstract":
            in_en_abstract = True
            in_cn_abstract = False
            i += 1
            continue

        if line.startswith("## "):
            if line == "## 0 引言" and not body_started:
                # Switch from the full-width front matter/overview figure to
                # the journal's two-column article body on the same page.
                section = doc.add_section(WD_SECTION.CONTINUOUS)
                configure_section(section, columns=2)
                body_started = True
            add_heading(doc, line[3:], 1)
            in_cn_abstract = in_en_abstract = False
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 2)
            i += 1
            continue

        if line.startswith("!["):
            m = re.search(r"\((.+)\)$", line)
            if m:
                image_path = ROOT / m.group(1)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                # The Normal style uses exact 13.5 pt leading. Inline pictures
                # would otherwise be clipped to that line box by LibreOffice.
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.add_run().add_break(WD_BREAK.PAGE)
                p.add_run().add_picture(str(image_path), width=Cm(16.7))
            i += 1
            continue

        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.right_indent = Cm(0.75)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            add_text_run(p, line[2:], size=8.5, font="楷体", color=(90, 90, 90))
            i += 1
            continue

        if in_cn_abstract:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.right_indent = Cm(0.75)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(15.6)
            add_text_run(p, "摘要：", size=9, bold=True, font="黑体")
            add_text_run(p, clean_inline(line), size=9, font="宋体")
            in_cn_abstract = False
            i += 1
            continue
        if in_en_abstract:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.right_indent = Cm(0.75)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(13.5)
            add_text_run(p, "Abstract: ", size=9, bold=True, font="Times New Roman")
            add_text_run(p, clean_inline(line), size=9, font="Times New Roman")
            in_en_abstract = False
            i += 1
            continue

        if line.startswith("**关键词：**") or line.startswith("**Keywords:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.right_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(2)
            label, value = line.split("**", 2)[1], line.split("**", 2)[2].strip()
            add_text_run(p, label, size=9, bold=True, font="黑体" if "关键" in label else "Times New Roman")
            add_text_run(p, value, size=9, font="宋体" if "关键" in label else "Times New Roman")
            i += 1
            continue

        if line.startswith("**中图分类号") or line.startswith("**GB/T 13745"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.right_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(1)
            add_text_run(p, clean_inline(line), size=9)
            i += 1
            continue

        if line.startswith("**图 ") or line.startswith("**Fig. ") or line.startswith("**表 "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(1)
            add_text_run(p, clean_inline(line), size=9, font="宋体" if not line.startswith("**Fig") else "Times New Roman")
            i += 1
            continue

        if re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(12)
            p.paragraph_format.space_after = Pt(1)
            add_text_run(p, clean_inline(line), size=8.5)
            i += 1
            continue

        if front_index == 0 and line.startswith("**Lifecycle-"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            add_text_run(p, clean_inline(line), size=12, bold=True, font="Times New Roman")
            front_index += 1
            i += 1
            continue
        if line.startswith("**作者"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_text_run(p, clean_inline(line), size=12, font="仿宋")
            i += 1
            continue
        if line.startswith("（作者单位"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            add_text_run(p, line, size=7.5)
            i += 1
            continue

        is_formula = ("=" in line and len(line) < 100 and not line.startswith("http"))
        add_body_paragraph(doc, line, first_indent=not is_formula, center=is_formula)
        i += 1

    props = doc.core_properties
    props.title = "含能材料制造多源风险证据的生命周期一致组合前治理方法（演示稿）"
    props.subject = "《兵工学报》版式演示；GB/T 13745—2009 主分类 5352045"
    props.keywords = "含能材料；证据融合；生命周期一致性；组合前治理"
    props.comments = "演示稿，非正式投稿；未伪造正式实验结果。"

    # Both sections share the same footer. Add the PAGE field once so a
    # continuous section break does not duplicate it (1, 2, 3 instead of
    # 11, 22, 33).
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p._element.clear_content()
    add_page_field(footer_p)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
